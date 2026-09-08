"""
Checks GST Command Support files attached to commentCommandSupport for the
existence of a "SID一覧" sheet and, if present, the validity of its 適用
（○/×）欄.

The validation logic below (Grid/Report, find_anchors, check_sidlist_column,
etc.) is copied in directly from check_gst_ox.py's macro-accurate validator
(reverse-engineered from the real VBA macro - modChkExec.ReadFromGST - that
actually consumes these files), rather than imported from that file, so this
module has no dependency on check_gst_ox.py at all. check_gst_ox.py itself
is untouched and still works standalone as its own CLI tool - this is a
copy, not a move.

【メッセージ書式】check_sid_list_sheet_parts系が返すmessageは、結果ログ（`check2_report`が
[GSTコマンド装備表]見出しの下へ並べる）の箇条書き1行としてそのまま使える文にする。
OK/NGはラベルではなく文末の肯定形／否定形（シートがある↔ない、データ不備がない↔ある）で
表し、NGの場合は続けて不備箇所（セル番地等）を書く。文言を変える場合は
`definition_file_check`／`matrix_check`側の文と揃えること（2026-09-03改修）。

【2026-09-02改修】unified_main.pyが同梱コピー（本フォルダ）をimportするため、判定条件が
`..\\gst_command_check_ref.py`（マクロ準拠版）と食い違わないよう、本ファイルの内容を
同ファイルと同一の実装へ差し替えた。以前の版は「SAE J1979／-2／-3 の各行を
"J1979を適用する場合"列の直前まで走査し、○×記号リストに含まれるか見るだけ」の簡易判定
だったため、マクロが実際に停止・誤読するケース（SID列の空欄による静かな打ち切り、
Description列と適用列の終端行不一致=MSG_GST_ROWCNT_ERR、通信プロトコルの○が1個でない、
数式・結合セル）を見逃していた。両ファイルは重複管理のため、改修時は必ず両方へ反映すること。

Different from the FS Matrix check (matrix_check.py): a missing attachment
is NOT an error here. If nothing is attached, check_gst_command_files()
simply returns OK with no message - the caller doesn't need to special-case
"no attachment" separately, it can always call this function and let it
decide there is nothing to check. When there IS at least one attachment,
every file whose name contains "CommandSupportGST" and has an Excel
extension gets checked for SID一覧 - files that don't match that pattern are
ignored, not flagged.
"""

import os

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SID_LIST_SHEET_NAME = "SID一覧"

# --- Copied from check_gst_ox.py: macro-side constants ------------------
SHEET_SIDLIST = SID_LIST_SHEET_NAME
OX_O = '○'                 # GBL_OX_O  ○ U+25CB
OX_X = '×'                 # GBL_OX_X  × U+00D7
SID_TITLE = 'Service ID(SID)'   # STR_SID_TITLE
DWG_REMOVE = ('37805-', '37806-', '28100-', '28101-',
              '5K800-', '5K801-', 'ZK800-', 'ZK801-')   # GST_DWG_REMOVE_SOFT
MAX_CLM = 256                   # NUM_MAX_CLM
XL_LAST_ROW = 1048576

# SID -> Service系シート名 (modChkExec.bas の Select Case strSID)
SID_TO_SHEET = {
    '$01': 'DataStream',
    '$02': 'FreezeFrame',
    '$06': 'TestResult',
    '$09': 'VehicleInformation',
    '$19-$04': 'FreezeFrame',
    '$22-$F4/$F5': 'DataStream',
    '$22-$F6': 'TestResult',
    '$22-$F8': 'VehicleInformation',
}
# Service系シートの列オフセット (startからの差分)
OFF_SUB1 = 0
OFF_DESCRIP = 3
OFF_06_DESCRIP = 6

PROTO_LABELS = ((2, 'SAE J1979'), (3, 'SAE J1979-2'), (4, 'SAE J1979-3'))


def is_command_support_gst_file(name):
    return "CommandSupportGST" in name and name.lower().endswith((".xlsx", ".xls"))


# --- Copied from check_gst_ox.py: VBA-equivalent sheet wrapper -----------
class Grid:
    """VBA の .Value / IsEmpty / End(xlDown) をそのまま再現するシートラッパ。"""

    def __init__(self, ws_val, ws_raw):
        self._v = ws_val    # data_only=True  : キャッシュ値
        self._r = ws_raw    # data_only=False : 数式
        self.title = ws_val.title
        self.max_row = max(ws_val.max_row or 1, ws_raw.max_row or 1)
        self.max_col = max(ws_val.max_column or 1, ws_raw.max_column or 1)
        # 結合セルの2番目以降は VBA でも Empty になる
        self.merged_followers = set()
        self.merged_all = set()
        for rng in ws_val.merged_cells.ranges:
            for row in range(rng.min_row, rng.max_row + 1):
                for col in range(rng.min_col, rng.max_col + 1):
                    self.merged_all.add((row, col))
                    if (row, col) != (rng.min_row, rng.min_col):
                        self.merged_followers.add((row, col))

    def formula(self, row, col):
        if row < 1 or col < 1:
            return None
        f = self._r.cell(row=row, column=col).value
        return f if isinstance(f, str) and f.startswith('=') else None

    def text(self, row, col):
        """VBA の .Value を文字列化したもの（未入力は ''）。トリムしない。"""
        if row < 1 or col < 1 or (row, col) in self.merged_followers:
            return ''
        v = self._v.cell(row=row, column=col).value
        if v is None:
            return ''
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v)

    def has_content(self, row, col):
        """End(xlDown) / IsEmpty が「空でない」と見なすか。

        数式セルは結果が "" でも Excel 的には空ではない点まで再現する。
        """
        if row < 1 or col < 1 or (row, col) in self.merged_followers:
            return False
        if self.formula(row, col) is not None:
            return True
        return self._v.cell(row=row, column=col).value is not None

    def end_down(self, row, col):
        """Range(row, col).End(xlDown).Row を再現。"""
        limit = self.max_row + 2
        if self.has_content(row + 1, col):
            r = row + 1
            while r <= limit and self.has_content(r + 1, col):
                r += 1
            return r
        r = row + 1
        while r <= limit:
            if self.has_content(r, col):
                return r
            r += 1
        return XL_LAST_ROW

    def a1(self, row, col):
        return '%s%d' % (get_column_letter(col), row)


class Report:
    def __init__(self):
        self.ng = []
        self.warn = []
        self.info = []

    def add_ng(self, where, msg):
        self.ng.append((where, msg))

    def add_warn(self, where, msg):
        self.warn.append((where, msg))

    def add_info(self, msg):
        self.info.append(msg)


def describe(value):
    """同形異字・空白混入を目視で判別できる形にする。"""
    if value == '':
        return '(空欄)'
    return '%r' % (value,)


def strip_dwg(raw):
    """GST_DWG_REMOVE_SOFT と改行除去を再現。"""
    s = raw.replace('\n', '')
    for pre in DWG_REMOVE:
        s = s.replace(pre, '')
    return s


def find_anchors(g, rep):
    """SIDデータ開始セルと DWG No.行を導出する。"""
    data_row = data_col = dwg_row = None
    scan_rows = min(g.max_row, 200)
    scan_cols = min(g.max_col, MAX_CLM)

    for row in range(1, scan_rows + 1):
        for col in range(1, scan_cols + 1):
            if g.text(row, col).strip() == SID_TITLE:
                data_row, data_col = row + 1, col
                break
        if data_row:
            break
    if data_row is None:
        rep.add_ng(SHEET_SIDLIST,
                   '「%s」が見つからず SIDデータ開始セルを特定できません。' % SID_TITLE)

    for row in range(1, scan_rows + 1):
        for col in range(1, scan_cols + 1):
            if g.text(row, col).replace('\n', '').strip() == 'SAE J1979':
                dwg_row = row - 2
                break
        if dwg_row:
            break
    if dwg_row is None:
        rep.add_ng(SHEET_SIDLIST,
                   '「SAE J1979」ラベルが見つからず DWG No.行を特定できません。')
    return data_row, data_col, dwg_row


def find_software_columns(g, data_col, dwg_row, target):
    """SID一覧のソフト列を抽出。

    マクロ (GstChk_DWG) と同じ判定にする: DWG No.行を右へ走査し、空セルで打ち切る。
    ○/× の有無では絞り込まない（適用欄が空のソフト列こそ検出したい対象なので）。
    ラベル用の結合セル（例: G10:J10）は結合セル判定で除外する。
    """
    limit = min(g.max_col, MAX_CLM)
    start = None
    for col in range(max(data_col + 1, 1), limit + 1):
        if (dwg_row, col) in g.merged_all:
            continue
        if g.text(dwg_row, col).strip():
            start = col
            break
    if start is None:
        return []

    out = []
    for col in range(start, limit + 1):
        raw = g.text(dwg_row, col)
        if not raw.strip():
            break                       # マクロもここでループを抜ける
        dwg = strip_dwg(raw)
        if target is None or strip_dwg(target) == dwg:
            out.append((col, dwg))
    return out


def find_dwg_column_on_sheet(g, dwg_row, target_dwg):
    """Service系シート側の DWG列 (マクロは各シートで再探索する)。"""
    for col in range(1, min(g.max_col, MAX_CLM) + 1):
        raw = g.text(dwg_row, col)
        if raw.strip() and strip_dwg(raw) == target_dwg:
            return col
    return None


def check_protocol_block(g, rep, where, dwg_row, col):
    """dwg_row+2..+4 (SAE J1979 / -2 / -3) の ○ が1つだけかを検証。

    マクロはこの3セルを「== ○ か否か」でしか見ないため書式エラーにはならない。
    ○ が0個/2個以上のときだけ BUBAN.csv の通信プロトコルが "ERR" になる。
    """
    marks = []
    n_o = 0
    for offset, label in PROTO_LABELS:
        v = g.text(dwg_row + offset, col).replace('\n', '')
        marks.append('%s=%s' % (label, v if v else '空欄'))
        if v == OX_O:
            n_o += 1
    if n_o != 1:
        rep.add_warn(where,
                     '%s:%s の ○ が %d 個です（%s）。1個でないと BUBAN.csv の通信'
                     'プロトコルが "ERR" になり、以降の検証が無警告で狂います。'
                     % (g.a1(dwg_row + 2, col), g.a1(dwg_row + 4, col),
                        n_o, ' / '.join(marks)))


def sidlist_end_rows(g, data_row, data_col, yryr_col):
    """SID一覧の終端行 (MSG_GST_ROWCNT_ERR の判定ロジック) を再現。

    Description列は +2、適用列は +3 で UDS ブロックの有無を見る非対称な実装。
    """
    desc_col = data_col + 1
    n1 = g.end_down(data_row, desc_col)
    if g.text(n1 + 2, desc_col) != '':
        n1 = g.end_down(n1, desc_col)
        n1 = g.end_down(n1, desc_col)
    n2 = g.end_down(data_row, yryr_col)
    if g.text(n2 + 3, yryr_col) != '':
        n2 = g.end_down(n2, yryr_col)
        n2 = g.end_down(n2, yryr_col)
    return n1, n2


def check_cell_hazards(g, rep, where, row, col):
    """数式・結合セルという「静かに壊れる」要因を報告。"""
    if g.formula(row, col) is not None:
        rep.add_warn(where,
                     '%s は数式です。キャッシュ値が古いとマクロは意図しない値を読みます。'
                     % g.a1(row, col))
    if (row, col) in g.merged_followers:
        rep.add_warn(where,
                     '%s は結合セルの2番目以降のため Empty として読まれます。'
                     % g.a1(row, col))


def check_sidlist_column(g, rep, data_row, data_col, dwg_row, col, dwg, verbose):
    """1つのソフト列について SID一覧を検証し、{SID: ○/×} と終端行を返す。"""
    where = '%s / %s (%s列)' % (SHEET_SIDLIST, dwg, get_column_letter(col))

    check_protocol_block(g, rep, where, dwg_row, col)

    # 適用欄が丸ごと空のソフト列（DWG No.だけ登録されている状態）
    if g.end_down(data_row, col) >= XL_LAST_ROW and not g.has_content(data_row, col):
        rep.add_ng(where,
                   '%s列に適用（○/×）が1件も入力されていません。ソフト一覧には'
                   '出てくるためこのDWGを選ぶと MSG_GST_ROWCNT_ERR で停止します。'
                   % get_column_letter(col))
        return {}, None

    n1, n2 = sidlist_end_rows(g, data_row, data_col, col)
    if n1 != n2:
        rep.add_ng(where,
                   'Description列(%s)の終端=%s行、適用列(%s)の終端=%s行 で不一致です'
                   '（MSG_GST_ROWCNT_ERR）。'
                   % (get_column_letter(data_col + 1), n1, get_column_letter(col), n2))
    end_row = min(n1, n2)
    if end_row >= XL_LAST_ROW:
        rep.add_ng(where, '終端行を特定できません（列全体が空の可能性）。')
        return {}, None

    if verbose:
        rep.add_info('%s: 適用欄の検証範囲 %s:%s'
                     % (where, g.a1(data_row, col), g.a1(end_row, col)))

    sid_map = {}
    row = data_row
    while row <= end_row:
        sid = g.text(row, data_col)
        # マクロの読み飛ばし / 打ち切り条件を再現
        if sid == '' and g.text(row + 1, data_col).strip() == SID_TITLE:
            row += 1
            continue
        if sid.strip() == SID_TITLE:
            row += 1
            continue
        if not g.has_content(row, data_col):
            rep.add_warn(where,
                         '%s の SID が空欄です。マクロはこの行で読み込みを打ち切り、'
                         '以降（〜%s行）を無警告で無視します。'
                         % (g.a1(row, data_col), end_row))
            break

        raw = g.text(row, col)
        if raw not in (OX_O, OX_X):
            rep.add_ng(where,
                       '%s (SID %s) が ○/× ではありません: %s'
                       % (g.a1(row, col), sid, describe(raw)))
        else:
            sid_map[sid] = raw
        check_cell_hazards(g, rep, where, row, col)
        row += 1

    return sid_map, end_row


def check_child_sheet(g, rep, sid, data_row, data_col,
                      dwg_row, dwg, verbose):
    """Service系シート (DataStream 等) の適用欄を検証。"""
    col = find_dwg_column_on_sheet(g, dwg_row, dwg)
    if col is None:
        rep.add_ng('%s [SID %s]' % (g.title, sid),
                   '%s行に DWG No. %s が見つかりません（MSG_GST_DWG_SW_NOTFOUND）。'
                   % (dwg_row, dwg))
        return

    where = '%s / %s (%s列) [SID %s]' % (g.title, dwg, get_column_letter(col), sid)

    # $06 系のみ Description の列が異なる
    is_06 = (sid == '$06') or (sid == '$22-$F6' and g.title == 'TestResult')
    desc_col = data_col + (OFF_06_DESCRIP if is_06 else OFF_DESCRIP)
    # SUB1 は $22 系 / $19-$04 のみ 1列ずれる
    sub1_col = data_col + OFF_SUB1 + (1 if ('$22' in sid or sid == '$19-$04') else 0)

    n1 = g.end_down(data_row, desc_col)
    n2 = g.end_down(data_row, col)
    if n1 != n2:
        if g.text(data_row + 1, desc_col) != '' or g.text(data_row + 1, col) != '':
            rep.add_ng(where,
                       'Description列(%s)の終端=%s行、適用列(%s)の終端=%s行 で不一致です'
                       '（MSG_GST_ROWCNT_ERR）。'
                       % (get_column_letter(desc_col), n1, get_column_letter(col), n2))
            return
        n1 = data_row      # マクロ側の「2行目が空なら1行だけ」の逃げ道
    end_row = n1
    if end_row >= XL_LAST_ROW:
        rep.add_warn(where, '終端行を特定できません。')
        return

    if verbose:
        rep.add_info('%s: 適用欄の検証範囲 %s:%s (Description=%s列)'
                     % (where, g.a1(data_row, col), g.a1(end_row, col),
                        get_column_letter(desc_col)))

    for row in range(data_row, end_row + 1):
        # マクロの打ち切り条件: Description/適用 が空、または SUB1 が "-"
        if not g.has_content(row, desc_col) or not g.has_content(row, col):
            if row < end_row:
                rep.add_warn(where,
                             '%s行で Description/適用 が空欄のためマクロは読み込みを'
                             '打ち切り、〜%s行を無警告で無視します。' % (row, end_row))
            break
        if g.text(row, sub1_col) == '-':
            break

        raw = g.text(row, col)
        if raw not in (OX_O, OX_X):
            rep.add_ng(where, '%s が ○/× ではありません: %s'
                       % (g.a1(row, col), describe(raw)))
        check_cell_hazards(g, rep, where, row, col)


def check_workbook(xlsx_path, dwg=None, sid_only=False):
    """Programmatic entry point - runs the anchor-derivation and 適用
    （○/×）欄 validation above, across every DWG/software column found (or
    just `dwg` if given) and, unless sid_only, cascaded into each SID's
    related Service sheet. Returns the resulting Report: rep.ng entries
    mean the macro would error out and stop; rep.warn entries mean it would
    run but silently produce a wrong result - callers should usually treat
    both as a real problem, not just rep.ng."""
    wb_val = load_workbook(xlsx_path, data_only=True)
    wb_raw = load_workbook(xlsx_path, data_only=False)
    grids = {n: Grid(wb_val[n], wb_raw[n]) for n in wb_val.sheetnames}

    rep = Report()

    if SHEET_SIDLIST not in grids:
        rep.add_ng(SHEET_SIDLIST, '「%s」シートがありません（MSG_GST_EQUIPSHEET_ERR）。' % SHEET_SIDLIST)
        return rep
    g = grids[SHEET_SIDLIST]

    data_row, data_col, dwg_row = find_anchors(g, rep)
    if not (data_row and dwg_row):
        return rep

    cols = find_software_columns(g, data_col, dwg_row, dwg)
    if not cols:
        if dwg:
            rep.add_ng(SHEET_SIDLIST,
                       'DWG No. "%s" は %d行に見つかりません（MSG_GST_DWG_SW_NOTFOUND と同じ状態）。'
                       % (dwg, dwg_row))
        else:
            rep.add_ng(SHEET_SIDLIST,
                       'ソフト列を検出できませんでした。DWG No.行の推定が誤っている可能性があります。')
        return rep

    for col, d in cols:
        sid_map, _end_row = check_sidlist_column(g, rep, data_row, data_col, dwg_row, col, d, verbose=False)
        if sid_only:
            continue
        for sid in sid_map:
            sheet = SID_TO_SHEET.get(sid)
            if sheet and sheet in grids:
                check_child_sheet(grids[sheet], rep, sid, data_row, data_col, dwg_row, d, verbose=False)

    return rep


# 適用欄の呼び方（②の結果ログ共通。`matrix_check.OX_FIELD_LABEL`と揃えること）。
OX_FIELD_LABEL = "適用（〇/×）欄"


def check_sid_list_sheet_parts(local_path):
    """Returns [(result, message)] - the GST Command Support file's report
    bullets, in the order they appear under the [GSTコマンド装備表] heading:

      1. シートの有無（ツールが想定している「SID一覧」シートがある／ない）
      2. その適用（〇/×）欄のデータ不備の有無（NGなら不備箇所を続けて書く）

    OK if SID_LIST_SHEET_NAME exists as a sheet in local_path and
    check_workbook(sid_only=True) reports no NG or WARN entries for it. A
    WARN there means the macro would run without stopping but silently
    produce a wrong result, so it counts as NG here too, not just a hard NG.
    sid_only=True keeps this scoped to SID一覧 itself, not the related
    Service sheets check_workbook can also cascade into.

    シートが無い／ファイルを開けない場合は適用欄を検査できないので、2行目は返さない
    （返る行数は1〜2行。呼び出し側はリストを走査して扱うこと）。"""
    file_name = os.path.basename(local_path)

    try:
        wb = load_workbook(local_path, data_only=True)
    except Exception as e:
        return [("NG", f"「{file_name}」を開けない: {e}")]

    if SID_LIST_SHEET_NAME not in wb.sheetnames:
        return [
            (
                "NG",
                f"ツールが想定している「{SID_LIST_SHEET_NAME}」シートがない。"
                f"（このファイルにあるシート: {', '.join(wb.sheetnames)}）",
            )
        ]

    parts = [("OK", f"ツールが想定している「{SID_LIST_SHEET_NAME}」シートがある。")]

    try:
        # sid_only=True: check SID一覧's own 適用（○/×）欄 only - skip
        # cascading into the related Service sheets (DataStream/
        # FreezeFrame/TestResult/VehicleInformation).
        rep = check_workbook(local_path, sid_only=True)
    except Exception as e:
        parts.append(
            ("NG", f"「{SID_LIST_SHEET_NAME}」シートの{OX_FIELD_LABEL}にデータ不備がないか確認できない。検証中にエラー: {e}")
        )
        return parts

    if rep.ng or rep.warn:
        problems = " / ".join(f"[{where}] {msg}" for where, msg in (rep.ng + rep.warn))
        parts.append(("NG", f"「{SID_LIST_SHEET_NAME}」シートの{OX_FIELD_LABEL}にデータ不備がある。{problems}"))
    else:
        parts.append(("OK", f"「{SID_LIST_SHEET_NAME}」シートの{OX_FIELD_LABEL}にデータ不備がない。"))

    return parts


def check_gst_command_files_parts(local_paths):
    """Runs check_sid_list_sheet_parts on every given local file path whose
    name matches the CommandSupportGST Excel-file pattern - others are
    ignored - and returns all of their bullets as one flat
    [(result, message)] list.

    添付が無い（またはパターンに一致するファイルが無い）場合は空リストを返す。これが
    「サイレントスキップ」で、結果ログには[GSTコマンド装備表]の見出しごと出さない
    （エラーではない）。ファイルが2つ以上ある場合だけ各行の先頭へファイル名を付ける。"""
    matching_paths = [path for path in local_paths if is_command_support_gst_file(os.path.basename(path))]

    parts = []
    for local_path in matching_paths:
        for result, message in check_sid_list_sheet_parts(local_path):
            if len(matching_paths) > 1:
                message = f"「{os.path.basename(local_path)}」: {message}"
            parts.append((result, message))
    return parts


def _combine(parts):
    """箇条書きの行リストを、SharePoint列へ書き戻す1件の(result, message)へまとめる
    （②`unified_result.combine_results`と同じ「1つでもNGなら最初のNGを返す」規則）。"""
    for result, message in parts:
        if result != "OK":
            return "NG", message
    return "OK", " / ".join(message for _, message in parts if message)


def check_sid_list_sheet(local_path):
    """Returns one combined (result, message) for local_path - the same
    checks as check_sid_list_sheet_parts, merged into a single result for
    callers that can only hold one (e.g. the SharePoint write-back)."""
    return _combine(check_sid_list_sheet_parts(local_path))


def check_gst_command_files(local_paths):
    """Returns one combined (result, message) for every matching file: OK
    with no message when there is nothing to check (the "skip silently"
    case, not an error), otherwise OK only if every matching file passed and
    NG with whichever bullet failed first."""
    return _combine(check_gst_command_files_parts(local_paths))
