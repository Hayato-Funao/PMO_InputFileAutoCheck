"""
matrix_check.py／gst_command_check.py／definition_file_check.py（②check_fs_matrixの
判定ロジック）を検証するための最小フィクスチャ（xlsx）を生成するモジュール
（2026-09-02新規。①(check_mao)の`テスト\\generate_fixtures.py`と同じ考え方で、
openpyxlで最小限のワークブックを組み立てる）。
"""

from openpyxl import Workbook

# matrix_check.py / gst_command_check.py が○×として許容する記号（両モジュール共通の定義を
# ここでも1つだけ使う。テストで「許容される記号」であることが分かればよいため、実データの
# 全表記ゆれを網羅する必要はない）。
OK_SYMBOL = "○"
NG_SYMBOL_INVALID = "△"  # ○×どちらでもない、想定外の記号


# matrix_check.detect_mark_columnsは「診断識別コード列(+0)〜SCS(+9)がデータ列」というマクロの
# 前提に従い、ソフト列（○/×列）を`診断識別コード列 + 9 + 1`より右だけで探す。また列全体の
# ○/×密度が5割以上の連続ブロックしか採らないため、最低5データ行を要する（NUM_IN_UDS_SCS
# = 9 / n_rows < 5 で打ち切りの条件）。フィクスチャはこの2条件を満たす必要がある。
MATRIX_MARK_COLUMN = 11     # 診断識別コード列(A=1) + 9 + 1
MATRIX_DATA_ROW_COUNT = 5   # detect_mark_columnsが密度判定を行う最小行数


def generate_matrix_fixture(path, sheet_name="FSマトリクス", table_symbol=OK_SYMBOL):
    """
    matrix_check.check_matrix_sheetが想定するFSマトリクスシートの最小構成を生成する
    （2026-09-02改修。matrix_checkをマクロ準拠版=matrix_check_ref.pyと同一実装へ差し替えた
    ことに伴い、そちらの判定条件を満たす構成へ作り替えた）。

    1〜3行目が部番／仕向け／ソフトのヘッダ3行（FsMatrix_Info.logが無い場合この3行は
    判定に使われないが、実データの形に合わせて置いている）、4行目に「診断識別コード」ラベル、
    5行目以降がMATRIX_DATA_ROW_COUNT行のデータ（A列に診断識別コード、
    MATRIX_MARK_COLUMN列に○）。

    `table_symbol`にOK_SYMBOL以外（例:NG_SYMBOL_INVALID）を渡すと、**最終データ行のみ**が
    その値になり、check_mark_blockがNGを返すデータになる。全行をその値にしないのは、
    ソフト列自体が○/×密度で検出される仕様上、列が丸ごと○×以外だと「○/×列を特定できません」
    という別のNGになってしまい、記号不備の検出を検証できないため。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    sheet.cell(row=1, column=1, value="部番")
    sheet.cell(row=1, column=MATRIX_MARK_COLUMN, value="37820-ABC-A000")
    sheet.cell(row=2, column=1, value="仕向け")
    sheet.cell(row=2, column=MATRIX_MARK_COLUMN, value="JPN")
    sheet.cell(row=3, column=1, value="ソフトウェア")
    sheet.cell(row=3, column=MATRIX_MARK_COLUMN, value="37805-ABC-A100")
    sheet.cell(row=4, column=1, value="診断識別コード")

    first_data_row = 5
    last_data_row = first_data_row + MATRIX_DATA_ROW_COUNT - 1
    for offset in range(MATRIX_DATA_ROW_COUNT):
        row = first_data_row + offset
        sheet.cell(row=row, column=1, value=f"D{offset + 1:03d}")
        symbol = table_symbol if row == last_data_row else OK_SYMBOL
        sheet.cell(row=row, column=MATRIX_MARK_COLUMN, value=symbol)

    workbook.save(path)


def generate_no_matrix_sheet_fixture(path):
    """matrix_check.MATRIX_SHEET_NAME_VARIANTSのいずれにも一致しないシート名のみを持つ
    ワークブックを生成する（check_matrix_sheetがNGになることの確認用）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "無関係なシート"
    sheet.cell(row=1, column=1, value="dummy")
    workbook.save(path)


def generate_gst_command_fixture(path, table_symbol=OK_SYMBOL):
    """
    gst_command_check.check_sid_list_sheetが想定する「SID一覧」シートの最小構成を生成する
    （2026-09-02改修。gst_command_checkをマクロ準拠版=gst_command_check_ref.pyと同一実装へ
    差し替えたことに伴い、そちらの判定条件を満たす構成へ作り替えた）。

    マクロ準拠版はアンカーを2つのラベルから導出する:
      * "Service ID(SID)"（STR_SID_TITLE）の1行下がSIDデータ開始行、その列がSID列
      * "SAE J1979"ラベル行の2行上がDWG No.行（＝ソフト列の見出し行）
    ソフト列はDWG No.行をSID列の右へ走査して最初に値が入っている列から連続する範囲で、
    その列の「DWG No.行+2〜+4」が通信プロトコル（SAE J1979 / -2 / -3）の○欄になる
    （○は必ず1個。0個や2個以上だとBUBAN.csvの通信プロトコルがERRになるため警告される）。
    SID列の各行はDescription列（SID列+1）と終端行が一致していなければならず
    （MSG_GST_ROWCNT_ERR）、SID列に空欄があるとマクロはそこで静かに打ち切る。

    上記を満たす最小構成として、C1にDWG No.、B3〜B5に通信プロトコルの3ラベル、
    C3に○（プロトコル欄）、A5に"Service ID(SID)"、6〜8行目をSIDデータ
    （A列=SID／B列=Description／C列=適用`table_symbol`）とする。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SID一覧"

    # DWG No.行（1行目）。SID列(A)とDescription列(B)は空、ソフト列(C)にDWG No.を置く
    # （find_software_columnsはSID列の右へ走査し、最初に値がある列をソフト列の先頭とみなす）。
    sheet.cell(row=1, column=3, value="37805-ABC-A100")

    # 通信プロトコル欄（DWG No.行+2〜+4）。○はSAE J1979の1個だけにする。
    sheet.cell(row=3, column=2, value="SAE J1979")
    sheet.cell(row=3, column=3, value=OK_SYMBOL)
    sheet.cell(row=4, column=2, value="SAE J1979-2")
    sheet.cell(row=5, column=2, value="SAE J1979-3")

    # SIDデータ開始行の導出元ラベル（この1行下=6行目からがデータ）。
    sheet.cell(row=5, column=1, value="Service ID(SID)")

    for offset, (sid, description) in enumerate((
        ("$01", "DataStream"),
        ("$02", "FreezeFrame"),
        ("$06", "TestResult"),
    )):
        row = 6 + offset
        sheet.cell(row=row, column=1, value=sid)
        sheet.cell(row=row, column=2, value=description)
        sheet.cell(row=row, column=3, value=table_symbol)

    workbook.save(path)


def generate_no_sid_list_sheet_fixture(path):
    """「SID一覧」シートを持たないワークブックを生成する
    （check_sid_list_sheetがNGになることの確認用）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "無関係なシート"
    sheet.cell(row=1, column=1, value="dummy")
    workbook.save(path)


def generate_master_definition_fixture(path, ok_sheet_has_defect=False, ng_sheet_has_defect=False, defect_row_count=1):
    """
    definition_file_check.SERVER_REFERENCE_FILE_PATHが指す「マスタ定義ファイル管理ツール」の
    ローカルコピーを模したワークブックを生成する。「元表」「fcl_list_o_FI」「fcl_list_x_FI」の
    3シートを持つ。

    fcl_list_o_FI／fcl_list_x_FIはそれぞれ「HILSコード」ヘッダの直後の行にデータを1行持つ。
    `*_has_defect=True`にすると、そのシートのデータ行がA列のみ値ありB/C列が空（データ不備）に
    なる。`defect_row_count`を増やすと不備行を3行目から連続でその数だけ作る（NGメッセージの
    セル番地列挙・件数の上限を検証する用途）。
    """
    workbook = Workbook()
    motohyou_sheet = workbook.active
    motohyou_sheet.title = "元表"
    motohyou_sheet.cell(row=1, column=1, value="項目A")
    motohyou_sheet.cell(row=1, column=2, value="項目B")
    motohyou_sheet.cell(row=2, column=1, value="値1")
    motohyou_sheet.cell(row=2, column=2, value="値2")

    for sheet_name, has_defect in (("fcl_list_o_FI", ok_sheet_has_defect), ("fcl_list_x_FI", ng_sheet_has_defect)):
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(row=1, column=1, value=sheet_name)
        sheet.cell(row=2, column=1, value="HILSコード")
        if has_defect:
            for offset in range(defect_row_count):
                # B・C列が空のまま（データ不備）
                sheet.cell(row=3 + offset, column=1, value=f"H{offset + 1:03d}")
        else:
            sheet.cell(row=3, column=1, value="H001")
            sheet.cell(row=3, column=2, value="設定値B")
            sheet.cell(row=3, column=3, value="設定値C")

    workbook.save(path)


def generate_sharepoint_reference_fixture(path, matches_motohyou=True):
    """
    definition_file_check.compare_motohyou_sheetsが比較する「もう1つの固定参照ファイル」
    （SharePoint上の`FI-FailSafeMatrix元表_SP.xlsx`相当）の「FI FSmatrix」シートを生成する。

    `matches_motohyou=True`なら`generate_master_definition_fixture`の「元表」シートと
    セル単位で一致する内容にする（OK判定の確認用）。Falseなら内容を変えて不一致にする
    （NG判定の確認用）。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FI FSmatrix"

    if matches_motohyou:
        sheet.cell(row=1, column=1, value="項目A")
        sheet.cell(row=1, column=2, value="項目B")
        sheet.cell(row=2, column=1, value="値1")
        sheet.cell(row=2, column=2, value="値2")
    else:
        sheet.cell(row=1, column=1, value="項目A")
        sheet.cell(row=1, column=2, value="項目B")
        sheet.cell(row=2, column=1, value="値1")
        sheet.cell(row=2, column=2, value="値2_不一致")

    workbook.save(path)
