"""判定ロジック・結果出力・案件フォルダ走査のテスト用に、CAN ID照合の最小フィクスチャ
（MAO・CANマトリクス・CANテーブル、および案件フォルダ構造）を生成するモジュール。

【2026-09-01改修】チェック対象をCANテーブルからCANマトリクスへ変更したことに伴い、
`generate_matrix_fixture`（マトリクス用フィクスチャ）を新規追加し、`generate_fixtures`・
`generate_case_folder_fixture`・`generate_fi_icm_case_folder_fixture`が生成する
「①の判定対象ファイル」をCANマトリクスへ差し替えた。CANテーブル用の生成関数
（`generate_table_fixture`等）は、`can_info.py`のCANテーブル関数（当面未使用）に対する
レガシーテストのために残している。

生成するCAN IDと期待される判定（内部仕様書4.5節。MAO＋CANマトリクスの2ソース版）の対応:
	100h(256) : MAO + マトリクス               → OK
	200h(512) : MAOのみ（マトリクス欠け）      → NG
	300h(768) : MAO + マトリクス               → OK
	400h(1024): MAOのみ（マトリクス欠け）      → NG
	500h(1280): マトリクスのみ（MAOに無い）    → －（判定対象外）
"""

import zipfile
from pathlib import Path

from openpyxl import Workbook

# MAOのCAN ID記載行を模したテキスト（実データの表記"CAN ID=074h"に合わせる）
MAO_FIXTURE_LINES = [
	"CANAA100TX?E001?CAN ID=100h Test Signal A?-?1110?io-comsfcn3oi100-osl-0?HEX?11?",
	"CANAA200TX?E002?CAN ID=200h Test Signal B?-?1110?io-comsfcn3oi200-osl-0?HEX?11?",
	"CANAA300RX?E003?CAN ID=300h Test Signal C?-?1110?io-comsfcn3ii300-osl-0?HEX?11?",
	"CANAA400TX?E004?CAN ID=400h Test Signal D?-?1110?io-comsfcn3oi400-osl-0?HEX?11?",
]


def generate_mao_fixture(mao_path):
	"""MAOのCAN ID記載行を模したテキストをcp932で書き出す。"""
	with open(mao_path, "w", encoding="cp932") as mao_file:
		for line in MAO_FIXTURE_LINES:
			mao_file.write(line + "\n")


def generate_table_fixture(table_path):
	"""
	CANテーブルのレイアウト（実データのR21ヘッダに合わせる。2026-08-31追補で内容判別・
	ヘッダ行動的検出に対応したため、ヘッダ行にも実際の見出し文字列を書き出す）を模したxlsxを生成する。

	R21がヘッダ（A列=Transmitter、B列=Receivers、C列=Signal Name、H列=ID、I列=ID-Format）、
	R22以降がデータ。CAN ID 100h・300h・500hを含む（500hはMAOに無いIDとして「－」判定の確認に使う）。
	"""
	workbook = Workbook()
	sheet = workbook.active

	sheet.cell(row=21, column=1, value="Transmitter")
	sheet.cell(row=21, column=2, value="Receivers")
	sheet.cell(row=21, column=3, value="Signal Name")
	sheet.cell(row=21, column=8, value="ID")
	sheet.cell(row=21, column=9, value="ID-Format")

	sheet.cell(row=22, column=1, value="ECU_A")
	sheet.cell(row=22, column=2, value="ECU_C")
	sheet.cell(row=22, column=3, value="SignalA")
	sheet.cell(row=22, column=8, value="100")
	sheet.cell(row=22, column=9, value="Standard")

	sheet.cell(row=23, column=1, value="ECU_A")
	sheet.cell(row=23, column=3, value="SignalC")
	sheet.cell(row=23, column=8, value="300")
	sheet.cell(row=23, column=9, value="Standard")

	sheet.cell(row=24, column=1, value="ECU_B")
	sheet.cell(row=24, column=2, value="ECU_C")
	sheet.cell(row=24, column=3, value="SignalE")
	sheet.cell(row=24, column=8, value="500")
	sheet.cell(row=24, column=9, value="Standard")

	workbook.save(table_path)


def generate_matrix_fixture(
	matrix_path,
	header_row=1,
	can_id_column=1,
	include_bus_summary_row=False,
	include_gateway_route_row=False,
):
	"""
	CANマトリクスのレイアウトを模したxlsxを生成する（2026-09-01新規）。

	旧実装（ヘッダR1〜R3・A列固定）はレイアウトのばらつきに弱いため、①の判定ロジックは
	見出し文字列`CAN ID (HEX)`からの動的検出に変更した（`can_info.find_matrix_header_location`）。
	本関数は`header_row`・`can_id_column`を引数で可変にすることで、そのばらつき（ヘッダ行位置・
	CAN ID列位置）を再現できるようにしている。

	CAN ID 100h・300h・500hを含む（500hはMAOに無いIDとして「－」判定の確認に使う）。
	`include_bus_summary_row=True`の場合、シート末尾にバス使用率集計行（CAN ID列に"F1"）を
	追加する（除外されることの確認用）。`include_gateway_route_row=True`の場合、ゲートウェイ
	経路表記（CAN ID列に"F1->F4,F5"）の行を追加する（同様に除外確認用）。

	引数:
		matrix_path: 出力先パス
		header_row: ヘッダ行番号（1始まり。旧仕様はR1〜R3のいずれか。機種によるばらつきの再現用）
		can_id_column: CAN ID (HEX)列の列番号（1始まり。旧仕様はA列＝1。機種によるばらつきの再現用）
		include_bus_summary_row: Trueならバス使用率集計行を追加する
		include_gateway_route_row: Trueならゲートウェイ経路表記の行を追加する
	"""
	workbook = Workbook()
	sheet = workbook.active

	sheet.cell(row=header_row, column=can_id_column, value="CAN ID (HEX)")

	data_row = header_row + 1
	sheet.cell(row=data_row, column=can_id_column, value="100")
	data_row += 1
	sheet.cell(row=data_row, column=can_id_column, value="300")
	data_row += 1
	sheet.cell(row=data_row, column=can_id_column, value="500")
	data_row += 1

	if include_gateway_route_row:
		# ゲートウェイ経路表記("F1->F4,F5"等)はCAN IDではないため読み飛ばされるべき
		sheet.cell(row=data_row, column=can_id_column, value="F1->F4,F5")
		data_row += 1

	if include_bus_summary_row:
		# シート末尾のバス使用率(%)集計行（バス名"F1"〜"F12"）はCAN IDではないため読み飛ばされるべき
		sheet.cell(row=data_row, column=can_id_column, value="F1")
		data_row += 1

	workbook.save(matrix_path)


def generate_fixtures(directory):
	"""
	指定フォルダにMAO・CANマトリクスの2フィクスチャファイルを生成する（単一マスク相当。
	抽出・突合ロジック単体のテスト用。2026-09-01改修：CANテーブルからCANマトリクスへ変更）。

	引数:
		directory: フィクスチャの出力先フォルダ（無ければ作成する）

	戻り値:
		{"mao": Path, "matrix": Path} の辞書
	"""
	directory = Path(directory)
	directory.mkdir(parents=True, exist_ok=True)

	paths = {
		"mao": directory / "fixture.mao",
		"matrix": directory / "fixture_Matrix.xlsx",
	}

	generate_mao_fixture(paths["mao"])
	generate_matrix_fixture(paths["matrix"])

	return paths


def generate_case_folder_fixture(directory):
	"""
	`case_scan.py`のテスト用に、案件フォルダ（`01_INPUT`配下）の最小構造を生成する。

	【2026-09-01改修】チェック対象をCANテーブルからCANマトリクスへ変更したことに伴い、
	採用対象・除外対象を反転した（旧: Matrix含む=除外、CANテーブル=採用 → 新: Matrix含む=採用、
	CANテーブル=除外decoy）。

	構造:
		01_INPUT/
		    01_HEX関連/
		        Mask1_HEX(A)/fixture.mao   … MAOが1個ある正常なマスク
		        Mask2_HEX(B)/               … MAOが無い（対象外。構造的NGではない）
		    02_CAN関連/
		        matrix_No1.xlsx              … ファイル名に"Matrix"を含むCANマトリクス（採用対象）
		        table.xlsx                   … CANテーブル形式のdecoy（除外対象。Matrix名を含まない）

	引数:
		directory: フィクスチャの出力先フォルダ（無ければ作成する）

	戻り値:
		{"input_root": Path, "mask1_dir": Path, "mask2_dir": Path,
		 "matrix_path": Path, "table_decoy_path": Path} の辞書
	"""
	directory = Path(directory)
	input_root = directory / "01_INPUT"
	hex_root = input_root / "01_HEX関連"
	can_root = input_root / "02_CAN関連"

	mask1_dir = hex_root / "Mask1_HEX(A)"
	mask2_dir = hex_root / "Mask2_HEX(B)"
	mask1_dir.mkdir(parents=True, exist_ok=True)
	mask2_dir.mkdir(parents=True, exist_ok=True)
	can_root.mkdir(parents=True, exist_ok=True)

	generate_mao_fixture(mask1_dir / "fixture.mao")

	matrix_path = can_root / "matrix_No1.xlsx"
	table_decoy_path = can_root / "table.xlsx"
	generate_matrix_fixture(matrix_path)
	generate_table_fixture(table_decoy_path)  # 除外されることの確認が目的（Matrix名を含まないため対象外）

	return {
		"input_root": input_root,
		"mask1_dir": mask1_dir,
		"mask2_dir": mask2_dir,
		"matrix_path": matrix_path,
		"table_decoy_path": table_decoy_path,
	}


def generate_fi_icm_case_folder_fixture(directory):
	"""
	`01_HEX関連`配下がFI/ICMサブフォルダに分かれ、それぞれの下に同名の`Mask1_HEX(A)`フォルダを
	持つ構成のフィクスチャを生成する（202606_005_01・202608_006_01で実データ確認済み。
	マスクラベルの相対パス化がFI/ICM間でラベルを重複させないことの検証用。2026-08-31追補）。

	戻り値:
		{"input_root": Path} の辞書
	"""
	directory = Path(directory)
	input_root = directory / "01_INPUT"
	hex_root = input_root / "01_HEX関連"
	can_root = input_root / "02_CAN関連"

	fi_mask1 = hex_root / "FI" / "Mask1_HEX(A)"
	icm_mask1 = hex_root / "ICM" / "Mask1_HEX(A)"
	fi_mask1.mkdir(parents=True, exist_ok=True)
	icm_mask1.mkdir(parents=True, exist_ok=True)
	can_root.mkdir(parents=True, exist_ok=True)

	generate_mao_fixture(fi_mask1 / "fi.mao")
	generate_mao_fixture(icm_mask1 / "icm.mao")
	generate_matrix_fixture(can_root / "matrix_No1.xlsx")

	return {"input_root": input_root}


def generate_multi_sheet_table_fixture(table_path):
	"""
	1ブックにTransmitter/Receiversヘッダを持つシートが複数あるCDCフォーマット類似のフィクスチャを
	生成する（3DAA(CDC)の実データ確認結果に基づく。2026-08-31追補）。

	"変更履歴"（ヘッダ無し。先頭シートがCANテーブル本体でないケースの再現）、"xFCAN1"
	（本体相当。シート名に"can"を含む）、"作業用"（ヘッダはあるがシート名に"can"/"communication"を
	含まない。変更後・削除リスト等の作業用シート相当）の3シートを持つ。
	`select_can_table_sheet`がシート名で"xFCAN1"を一意に選べることの検証用。
	"""
	workbook = Workbook()

	change_history_sheet = workbook.active
	change_history_sheet.title = "変更履歴"
	change_history_sheet.cell(row=1, column=1, value="変更日")

	main_sheet = workbook.create_sheet("xFCAN1")
	main_sheet.cell(row=1, column=1, value="Transmitter")
	main_sheet.cell(row=1, column=2, value="Receivers")
	main_sheet.cell(row=1, column=3, value="Signal Name")
	main_sheet.cell(row=1, column=8, value="ID")
	main_sheet.cell(row=1, column=9, value="ID-Format")
	main_sheet.cell(row=2, column=1, value="ECU_A")
	main_sheet.cell(row=2, column=8, value="100")
	main_sheet.cell(row=2, column=9, value="Standard")

	work_sheet = workbook.create_sheet("作業用")
	work_sheet.cell(row=1, column=1, value="Transmitter")
	work_sheet.cell(row=1, column=2, value="Receivers")
	work_sheet.cell(row=1, column=3, value="Signal Name")
	work_sheet.cell(row=1, column=8, value="ID")
	work_sheet.cell(row=1, column=9, value="ID-Format")

	workbook.save(table_path)
	return table_path


def generate_ambiguous_multi_sheet_table_fixture(table_path):
	"""
	シート名に"CAN"を含むTransmitter/Receiversヘッダシートが2つあり、`select_can_table_sheet`が
	一意に選べない（=判定不可とすべき）ケースのフィクスチャを生成する（2026-08-31追補）。
	"""
	workbook = Workbook()

	sheet_a = workbook.active
	sheet_a.title = "CAN_A"
	sheet_a.cell(row=1, column=1, value="Transmitter")
	sheet_a.cell(row=1, column=2, value="Receivers")

	sheet_b = workbook.create_sheet("CAN_B")
	sheet_b.cell(row=1, column=1, value="Transmitter")
	sheet_b.cell(row=1, column=2, value="Receivers")

	workbook.save(table_path)
	return table_path


def generate_3daa_cdc_style_matrix_fixture(matrix_path):
	"""
	3DAA(CDC)形式のCANマトリクスを模したフィクスチャを生成する（2026-09-01追補）。

	実データ202608_006_01の`Matrix_3DAA_FHEV_PADAS(CDC)_F_ALL_04.03.00.xlsx`で確認された、
	ファイル名は"Matrix"を含むが内部の列構成が`CAN ID (HEX)`単独列ではなく、`ID`・`Transmitter`・
	`Receivers`というCANテーブルに近い列構成（列位置はB=ID/C=Transmitter/D=Receivers）に
	なっているケースの再現用。`find_matrix_header_location`のフォールバック検出（Transmitter/
	Receiversヘッダ+ID列）の検証に使う。
	"""
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Matrix"

	sheet.cell(row=3, column=2, value="ID")
	sheet.cell(row=3, column=3, value="Transmitter")
	sheet.cell(row=3, column=4, value="Receivers")

	sheet.cell(row=7, column=2, value="100")
	sheet.cell(row=7, column=3, value="ENG")
	sheet.cell(row=7, column=4, value="ICM")

	sheet.cell(row=8, column=2, value="300")
	sheet.cell(row=8, column=3, value="ICM")
	sheet.cell(row=8, column=4, value="ENG")

	workbook.save(matrix_path)
	return matrix_path


def generate_multi_sheet_matrix_fixture(matrix_path):
	"""
	1ブックに`CAN ID (HEX)`ヘッダを持つシートが複数あるフィクスチャを生成する（2026-09-01新規）。

	CANマトリクスには、CANテーブルの3DAA(CDC)のようなシート名の絞り込みキーワードが実データで
	確認されていないため、`select_matrix_sheet`は複数候補時に絞り込みを行わずNone（判定不可）を
	返す仕様になっている。本フィクスチャはその挙動の検証用。
	"""
	workbook = Workbook()

	sheet_a = workbook.active
	sheet_a.title = "Sheet_A"
	sheet_a.cell(row=1, column=1, value="CAN ID (HEX)")
	sheet_a.cell(row=2, column=1, value="100")

	sheet_b = workbook.create_sheet("Sheet_B")
	sheet_b.cell(row=1, column=1, value="CAN ID (HEX)")
	sheet_b.cell(row=2, column=1, value="200")

	workbook.save(matrix_path)
	return matrix_path


def zip_directory_contents(source_directory, zip_path):
	"""source_directory配下の内容をまとめてzip_pathへzip化する（zip解凍テスト用）。"""
	source_directory = Path(source_directory)
	zip_path = Path(zip_path)

	with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
		for path in source_directory.rglob("*"):
			if path.is_file():
				zip_file.write(path, path.relative_to(source_directory))

	return zip_path
